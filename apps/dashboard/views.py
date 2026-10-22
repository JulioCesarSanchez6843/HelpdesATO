from datetime import timedelta
from io import BytesIO

from django.db.models import (
    Avg, Count, DurationField, ExpressionWrapper, F, OuterRef, Q, Subquery,
)
from django.db.models.functions import TruncDate, TruncMonth
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import dateformat, timezone
from django.utils.dateparse import parse_date
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.accounts.decorators import admin_required
from apps.tickets.models import Historial, Ticket

DURACION_CIERRE = ExpressionWrapper(
    F("cerrado_en") - F("creado_en"), output_field=DurationField()
)
FILTRO_CERRADOS_CON_FECHA = Q(estado=Ticket.Estado.CERRADO, cerrado_en__isnull=False)
MESES = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]


def _horas(duracion):
    return round(duracion.total_seconds() / 3600, 1) if duracion else None


def _fecha_get(request, nombre):
    """Lee una fecha YYYY-MM-DD de la URL; si viene mal escrita, la ignora."""
    try:
        return parse_date(request.GET.get(nombre, "").strip())
    except ValueError:
        return None

def _tickets_alcance(request):
    tickets = Ticket.objects.all()
    if request.user.es_superadmin:
        return tickets
    return tickets.filter(sucursal_id=request.user.sucursal_id)

def _resumen_por(tickets, campo, sin_valor):
    """Agrupa los tickets por `campo` (ej. 'area__nombre') y calcula, por grupo:
    total, pendientes, cerrados, % cerrados y tiempo promedio de cierre."""
    filas = (
        tickets.order_by()
        .values(campo)
        .annotate(
            total=Count("id"),
            cerrados=Count("id", filter=Q(estado=Ticket.Estado.CERRADO)),
            promedio=Avg(DURACION_CIERRE, filter=FILTRO_CERRADOS_CON_FECHA),
        )
        .order_by("-total", campo)
    )
    resultado = []
    for fila in filas:
        total, cerrados = fila["total"], fila["cerrados"]
        resultado.append({
            "nombre": fila[campo] or sin_valor,
            "total": total,
            "pendientes": total - cerrados,
            "cerrados": cerrados,
            "porcentaje": round(cerrados / total * 100) if total else 0,
            "horas": _horas(fila["promedio"]),
        })
    return resultado


def _grafica_barras(filas):
    return {
        "labels": [f["nombre"] for f in filas],
        "pendientes": [f["pendientes"] for f in filas],
        "cerrados": [f["cerrados"] for f in filas],
    }

@admin_required
def reportes(request):
    desde = _fecha_get(request, "desde")
    hasta = _fecha_get(request, "hasta")

    base = _tickets_alcance(request)
    tickets = base
    if desde:
        tickets = tickets.filter(creado_en__date__gte=desde)
    if hasta:
        tickets = tickets.filter(creado_en__date__lte=hasta)

    total = tickets.count()
    abiertos = tickets.filter(estado=Ticket.Estado.ABIERTO).count()
    en_proceso = tickets.filter(estado=Ticket.Estado.EN_PROCESO).count()
    por_confirmar = tickets.filter(estado=Ticket.Estado.POR_CONFIRMAR).count()
    cerrados = tickets.filter(estado=Ticket.Estado.CERRADO).count()

    promedio = tickets.filter(FILTRO_CERRADOS_CON_FECHA).aggregate(
        promedio=Avg(DURACION_CIERRE)
    )["promedio"]
    tiempo_promedio_horas = _horas(promedio)

    porcentaje_resueltos = round((cerrados / total) * 100) if total else 0

    hoy = timezone.localdate()
    dias = [hoy - timedelta(days=i) for i in range(6, -1, -1)]
    por_dia = dict(
        base.filter(creado_en__date__gte=dias[0])
        .annotate(dia=TruncDate("creado_en"))
        .order_by()
        .values_list("dia")
        .annotate(total=Count("id"))
    )

        
    primer_mes = hoy.replace(day=1)
    meses = []
    for _ in range(6):
        meses.append(primer_mes)
        primer_mes = (primer_mes - timedelta(days=1)).replace(day=1)
    meses.reverse()

    def _por_mes(campo):
        datos = (
            base.filter(**{f"{campo}__date__gte": meses[0]})
            .annotate(mes=TruncMonth(campo))
            .order_by()
            .values_list("mes")
            .annotate(total=Count("id"))
        )
        return {m.date() if hasattr(m, "date") else m: n for m, n in datos}

    creados_mes = _por_mes("creado_en")
    cerrados_mes = _por_mes("cerrado_en")

    por_categoria = _resumen_por(tickets, "categoria__nombre", "Sin categoría")
    por_area = _resumen_por(tickets, "area__nombre", "Sin área")
    por_sucursal = _resumen_por(tickets, "sucursal__nombre", "Sin sucursal")

    top_solicitantes = list(
        tickets.order_by()
        .values("creado_por__nombre", "creado_por__area__nombre")
        .annotate(total=Count("id"))
        .order_by("-total", "creado_por__nombre")[:5]
    )

    graficas = {
        "actividad": {
            "labels": [dateformat.format(d, "d-b") for d in dias],
            "data": [por_dia.get(d, 0) for d in dias],
        },
        "distribucion": {
            "labels": ["Abiertos", "En proceso", "Por confirmar", "Cerrados"],
            "data": [abiertos, en_proceso, por_confirmar, cerrados],
        },
        "categorias": _grafica_barras(por_categoria),
        "areas": _grafica_barras(por_area),
        "mensual": {
            "labels": [f"{MESES[m.month - 1]} {m.year}" for m in meses],
            "creados": [creados_mes.get(m, 0) for m in meses],
            "cerrados": [cerrados_mes.get(m, 0) for m in meses],
        },
    }

    return render(request, "dashboard/reportes.html", {
        "total": total,
        "abiertos": abiertos,
        "en_proceso": en_proceso,
        "por_confirmar": por_confirmar,
        "cerrados": cerrados,
        "tiempo_promedio_horas": tiempo_promedio_horas,
        "porcentaje_resueltos": porcentaje_resueltos,
        "graficas": graficas,
        "por_categoria": por_categoria,
        "por_area": por_area,
        "por_sucursal": por_sucursal,
        "top_solicitantes": top_solicitantes,
        "desde": desde.isoformat() if desde else "",
        "hasta": hasta.isoformat() if hasta else "",
    })

def _local(fecha):
    """Excel no admite fechas con zona horaria: se pasa a hora local sin tzinfo."""
    return timezone.localtime(fecha).replace(tzinfo=None) if fecha else None


def _celda_segura(valor):
    """Quita caracteres que Excel no acepta."""
    if isinstance(valor, str):
        return ILLEGAL_CHARACTERS_RE.sub("", valor)
    return valor


def _escribir(hoja, fila, columna, valor, formato=None):
    celda = hoja.cell(row=fila, column=columna, value=_celda_segura(valor))
    if isinstance(valor, str):
        celda.data_type = "s"
        if valor[:1] in ("=", "+", "-", "@"):
            celda.quotePrefix = True
    if formato:
        celda.number_format = formato
    return celda


def _estilizar_encabezado(hoja, fila, columnas):
    relleno = PatternFill("solid", fgColor="1F4E78")
    for col in range(1, columnas + 1):
        celda = hoja.cell(row=fila, column=col)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = relleno
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _ultima_accion(texto):
    """Subconsulta: último movimiento del historial cuyo texto contiene `texto`."""
    return (
        Historial.objects.filter(
            ticket=OuterRef("pk"),
            tipo=Historial.Tipo.CAMBIO_ESTADO,
            contenido__contains=texto,
        )
        .order_by("-fecha", "-id")
    )


@admin_required
def exportar_resueltos(request):
    """Descarga en Excel los tickets resueltos (cerrados) con todos sus datos.
    Se puede acotar por fecha de cierre con ?desde=AAAA-MM-DD&hasta=AAAA-MM-DD."""
    desde = _fecha_get(request, "desde")
    hasta = _fecha_get(request, "hasta")

    resueltos_por = _ultima_accion("a «Por confirmar»")
    cierres = _ultima_accion("a «Cerrado»")

    tickets = (
        _tickets_alcance(request).filter(estado=Ticket.Estado.CERRADO)
        .select_related("area", "categoria", "creado_por", "creado_por__area", "creado_por__puesto")
        .annotate(
            resuelto_por=Subquery(resueltos_por.values("usuario__nombre")[:1]),
            texto_cierre=Subquery(cierres.values("contenido")[:1]),
            num_adjuntos=Count("adjuntos"),
        )
        .order_by("-cerrado_en", "-id")
    )
    if desde:
        tickets = tickets.filter(cerrado_en__date__gte=desde)
    if hasta:
        tickets = tickets.filter(cerrado_en__date__lte=hasta)

    wb = Workbook()
    hoja = wb.active
    hoja.title = "Tickets resueltos"

    encabezados = [
        "ID", "Título", "Descripción", "Nº de serie / Service Tag",
        "Área (departamento)", "Categoría", "Estado",
        "Solicitante", "Correo del solicitante", "Área del solicitante", "Puesto del solicitante",
        "Fecha de creación", "Fecha de resolución", "Fecha de cierre",
        "Horas hasta resolver", "Horas hasta el cierre",
        "Resuelto por", "Forma de cierre", "Nota de resolución", "Adjuntos",
    ]
    hoja.append(encabezados)
    _estilizar_encabezado(hoja, 1, len(encabezados))

    formato_fecha = "dd/mm/yyyy hh:mm"
    resumen_area, resumen_categoria = {}, {}
    suma_horas, con_horas = 0.0, 0

    for fila, t in enumerate(tickets, start=2):
        horas_resolver = (
            round((t.resuelto_en - t.creado_en).total_seconds() / 3600, 1)
            if t.resuelto_en else None
        )
        horas_cierre = (
            round((t.cerrado_en - t.creado_en).total_seconds() / 3600, 1)
            if t.cerrado_en else None
        )
        if horas_cierre is not None:
            suma_horas += horas_cierre
            con_horas += 1

        if t.texto_cierre and "administrador" in t.texto_cierre:
            forma_cierre = "Cerrado por el administrador"
        elif t.texto_cierre:
            forma_cierre = "Confirmado por el usuario"
        else:
            forma_cierre = ""

        solicitante = t.creado_por
        valores = [
            t.pk, t.titulo, t.descripcion, t.numero_serie,
            t.area.nombre if t.area else "Sin área",
            t.categoria.nombre if t.categoria else "Sin categoría",
            t.get_estado_display(),
            solicitante.nombre, solicitante.correo,
            solicitante.area.nombre if solicitante.area else "",
            solicitante.puesto.nombre if solicitante.puesto else "",
            _local(t.creado_en), _local(t.resuelto_en), _local(t.cerrado_en),
            horas_resolver, horas_cierre,
            t.resuelto_por or "", forma_cierre, t.nota_resolucion, t.num_adjuntos,
        ]
        for col, valor in enumerate(valores, start=1):
            formato = formato_fecha if col in (12, 13, 14) else None
            celda = _escribir(hoja, fila, col, valor, formato)
            if col in (2, 3, 19):
                celda.alignment = Alignment(wrap_text=True, vertical="top")
            else:
                celda.alignment = Alignment(vertical="top")

        resumen_area[valores[4]] = resumen_area.get(valores[4], 0) + 1
        resumen_categoria[valores[5]] = resumen_categoria.get(valores[5], 0) + 1

    anchos = [8, 32, 50, 22, 20, 18, 12, 24, 28, 20, 20, 18, 18, 18, 12, 12, 22, 28, 50, 10]
    for i, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(i)].width = ancho
    hoja.row_dimensions[1].height = 32
    hoja.freeze_panes = "C2"
    hoja.auto_filter.ref = hoja.dimensions

    resumen = wb.create_sheet("Resumen")
    total = sum(resumen_area.values())
    resumen.append(["Resumen de tickets resueltos"])
    resumen["A1"].font = Font(bold=True, size=14)
    resumen.append(["Generado", _local(timezone.now())])
    resumen.cell(row=2, column=2).number_format = formato_fecha
    resumen.append(["Cierre desde", desde.isoformat() if desde else "Sin límite"])
    resumen.append(["Cierre hasta", hasta.isoformat() if hasta else "Sin límite"])
    resumen.append(["Total de tickets", total])
    resumen.append([
        "Horas promedio hasta el cierre",
        round(suma_horas / con_horas, 1) if con_horas else "—",
    ])
    for titulo, datos in (
        ("Por área (departamento)", resumen_area),
        ("Por categoría", resumen_categoria),
    ):
        resumen.append([])
        resumen.append([titulo, "Tickets"])
        _estilizar_encabezado(resumen, resumen.max_row, 2)
        for nombre, cantidad in sorted(datos.items(), key=lambda x: (-x[1], x[0])):
            resumen.append([_celda_segura(nombre), cantidad])
    resumen.column_dimensions["A"].width = 34
    resumen.column_dimensions["B"].width = 22

    salida = BytesIO()
    wb.save(salida)

    nombre = f"tickets_resueltos_{timezone.localdate().isoformat()}.xlsx"
    respuesta = HttpResponse(
        salida.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    respuesta["Content-Disposition"] = f'attachment; filename="{nombre}"'
    return respuesta