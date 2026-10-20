from datetime import timedelta

from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from apps.accounts.models import Usuario
from apps.catalog.models import Area
from apps.tickets.models import Ticket


class ReportesTests(TestCase):
    def setUp(self):
        self.area = Area.objects.create(nombre="Ventas")
        self.admin = Usuario.objects.create_user("a@t.com", "Admin", "pw-segura-123", rol="ADMIN")
        self.user = Usuario.objects.create_user("u@t.com", "User", "pw-segura-123")

    def test_usuario_normal_no_entra(self):
        self.client.force_login(self.user)
        r = self.client.get(reverse("reportes"))
        self.assertRedirects(r, reverse("lista_tickets"))

    def test_sin_tickets_no_falla(self):
        self.client.force_login(self.admin)
        r = self.client.get(reverse("reportes"))
        self.assertEqual(r.status_code, 200)
        self.assertIsNone(r.context["tiempo_promedio_horas"])
        self.assertEqual(r.context["porcentaje_resueltos"], 0)
        self.assertEqual(r.context["graficas"]["actividad"]["data"], [0] * 7)

    def test_kpis_y_actividad_reales(self):
        ahora = timezone.now()
        t = Ticket.objects.create(titulo="a", descripcion="d", area=self.area, creado_por=self.user)
        Ticket.objects.create(titulo="b", descripcion="d", area=self.area, creado_por=self.user)
        cerrado = Ticket.objects.create(
            titulo="c", descripcion="d", area=self.area, creado_por=self.user,
            estado="CERRADO", cerrado_en=ahora + timedelta(hours=4))
        Ticket.objects.filter(pk=cerrado.pk).update(creado_en=ahora)
        self.client.force_login(self.admin)
        ctx = self.client.get(reverse("reportes")).context
        self.assertEqual((ctx["total"], ctx["abiertos"], ctx["cerrados"]), (3, 2, 1))
        self.assertEqual(ctx["porcentaje_resueltos"], 33)
        self.assertEqual(ctx["tiempo_promedio_horas"], 4.0)
        self.assertEqual(sum(ctx["graficas"]["actividad"]["data"]), 3)
        self.assertEqual(ctx["graficas"]["actividad"]["data"][-1], 3)
        

class EstadisticasYExcelTests(TestCase):
    def setUp(self):
        from apps.catalog.models import Categoria, Puesto
        self.red = Categoria.objects.create(nombre="Red")
        self.hw = Categoria.objects.create(nombre="Hardware")
        self.sistemas = Area.objects.create(nombre="Sistemas")
        self.ventas = Area.objects.create(nombre="Ventas")
        self.puesto = Puesto.objects.create(nombre="Analista")
        self.admin = Usuario.objects.create_user("a@t.com", "Admin", "pw-segura-123", rol="ADMIN")
        self.user = Usuario.objects.create_user(
            "u@t.com", "Usuario Uno", "pw-segura-123", area=self.ventas, puesto=self.puesto)

    def _ticket(self, titulo, area, categoria, estado="ABIERTO", horas=None, **extra):
        t = Ticket.objects.create(
            titulo=titulo, descripcion="desc", area=area, categoria=categoria,
            creado_por=self.user, estado=estado, **extra)
        if horas is not None:
            ahora = timezone.now()
            Ticket.objects.filter(pk=t.pk).update(
                creado_en=ahora - timedelta(hours=horas),
                resuelto_en=ahora - timedelta(minutes=30),
                cerrado_en=ahora)
            t.refresh_from_db()
        return t

    def test_estadisticas_por_categoria_y_area(self):
        self._ticket("a", self.sistemas, self.red)
        self._ticket("b", self.sistemas, self.red, estado="CERRADO", horas=4)
        self._ticket("c", self.ventas, self.hw, estado="CERRADO", horas=2)
        self._ticket("d", None, None)
        self.client.force_login(self.admin)
        ctx = self.client.get(reverse("reportes")).context

        cats = {f["nombre"]: f for f in ctx["por_categoria"]}
        self.assertEqual(cats["Red"]["total"], 2)
        self.assertEqual(cats["Red"]["pendientes"], 1)
        self.assertEqual(cats["Red"]["cerrados"], 1)
        self.assertEqual(cats["Red"]["porcentaje"], 50)
        self.assertEqual(cats["Red"]["horas"], 4.0)
        self.assertEqual(cats["Hardware"]["horas"], 2.0)
        self.assertEqual(cats["Sin categoría"]["total"], 1)

        areas = {f["nombre"]: f for f in ctx["por_area"]}
        self.assertEqual(areas["Sistemas"]["total"], 2)
        self.assertEqual(areas["Ventas"]["cerrados"], 1)
        self.assertEqual(areas["Sin área"]["total"], 1)

        g = ctx["graficas"]
        self.assertEqual(sum(g["categorias"]["pendientes"]) + sum(g["categorias"]["cerrados"]), 4)
        self.assertEqual(sum(g["mensual"]["creados"]), 4)
        self.assertEqual(sum(g["mensual"]["cerrados"]), 2)
        self.assertEqual(ctx["top_solicitantes"][0]["total"], 4)

    def test_filtro_de_fechas_y_fecha_invalida(self):
        vieja = self._ticket("vieja", self.sistemas, self.red)
        Ticket.objects.filter(pk=vieja.pk).update(creado_en=timezone.now() - timedelta(days=60))
        self._ticket("nueva", self.sistemas, self.red)
        self.client.force_login(self.admin)
        desde = (timezone.localdate() - timedelta(days=5)).isoformat()
        self.assertEqual(self.client.get(reverse("reportes"), {"desde": desde}).context["total"], 1)
        # Una fecha mal escrita se ignora en lugar de romper la página.
        r = self.client.get(reverse("reportes"), {"desde": "no-es-fecha", "hasta": "2026-99-99"})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.context["total"], 2)

    def test_excel_solo_cerrados_con_sus_datos(self):
        from io import BytesIO
        from openpyxl import load_workbook
        from apps.tickets.models import Historial

        abierto = self._ticket("abierto", self.sistemas, self.red)
        cerrado = self._ticket(
            "cerrado", self.sistemas, self.red, estado="CERRADO", horas=5,
            numero_serie="ABC123", nota_resolucion="Se cambió el cable")
        Historial.objects.create(
            ticket=cerrado, usuario=self.admin, tipo="CAMBIO_ESTADO",
            contenido="Estado cambiado de «En proceso» a «Por confirmar». Nota de resolución: x")
        Historial.objects.create(
            ticket=cerrado, usuario=self.user, tipo="CAMBIO_ESTADO",
            contenido="Estado cambiado de «Por confirmar» a «Cerrado». El usuario confirmó que el problema quedó resuelto.")

        self.client.force_login(self.admin)
        r = self.client.get(reverse("exportar_resueltos"))
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r["Content-Type"])
        self.assertIn("attachment", r["Content-Disposition"])

        hoja = load_workbook(BytesIO(r.content))["Tickets resueltos"]
        filas = list(hoja.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(filas), 1)
        enc = [c.value for c in hoja[1]]
        fila = dict(zip(enc, filas[0]))
        self.assertEqual(fila["ID"], cerrado.pk)
        self.assertEqual(fila["Área (departamento)"], "Sistemas")
        self.assertEqual(fila["Categoría"], "Red")
        self.assertEqual(fila["Solicitante"], "Usuario Uno")
        self.assertEqual(fila["Área del solicitante"], "Ventas")
        self.assertEqual(fila["Puesto del solicitante"], "Analista")
        self.assertEqual(fila["Nº de serie / Service Tag"], "ABC123")
        self.assertEqual(fila["Horas hasta el cierre"], 5.0)
        self.assertEqual(fila["Resuelto por"], "Admin")
        self.assertEqual(fila["Forma de cierre"], "Confirmado por el usuario")
        self.assertEqual(fila["Nota de resolución"], "Se cambió el cable")
        self.assertNotIn(abierto.pk, [f[0] for f in filas])

    def test_excel_filtra_por_fecha_de_cierre(self):
        from io import BytesIO
        from openpyxl import load_workbook
        viejo = self._ticket("viejo", self.sistemas, self.red, estado="CERRADO", horas=1)
        Ticket.objects.filter(pk=viejo.pk).update(cerrado_en=timezone.now() - timedelta(days=40))
        self._ticket("reciente", self.sistemas, self.red, estado="CERRADO", horas=1)
        self.client.force_login(self.admin)
        desde = (timezone.localdate() - timedelta(days=7)).isoformat()
        r = self.client.get(reverse("exportar_resueltos"), {"desde": desde})
        hoja = load_workbook(BytesIO(r.content))["Tickets resueltos"]
        self.assertEqual(hoja.max_row, 2)

    def test_excel_no_ejecuta_formulas_ni_rompe_con_caracteres_raros(self):
        from io import BytesIO
        from openpyxl import load_workbook
        self._ticket("=HYPERLINK(\"http://malo\",\"x\")", self.sistemas, self.red,
                     estado="CERRADO", horas=1, nota_resolucion="+cmd\x07ok")
        self.client.force_login(self.admin)
        r = self.client.get(reverse("exportar_resueltos"))
        self.assertEqual(r.status_code, 200)
        hoja = load_workbook(BytesIO(r.content))["Tickets resueltos"]
        celda = hoja["B2"]
        self.assertEqual(celda.data_type, "s")
        self.assertTrue(celda.value.startswith("="))
        self.assertEqual(hoja["S2"].value, "+cmdok")

    def test_excel_vacio_y_permisos(self):
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(reverse("exportar_resueltos")).status_code, 200)
        self.client.force_login(self.user)
        r = self.client.get(reverse("exportar_resueltos"))
        self.assertRedirects(r, reverse("lista_tickets"))
        self.client.logout()
        r = self.client.get(reverse("exportar_resueltos"))
        self.assertEqual(r.status_code, 302)